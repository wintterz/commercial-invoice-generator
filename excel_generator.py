"""Render an Invoice to an .xlsx file using openpyxl."""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from invoice_models import Invoice

ACCENT = "1F3864"
LIGHT = "EAF0FA"

_thin = Side(style="thin", color="CCCCCC")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def invoice_to_xlsx(invoice: Invoice, path: str) -> str:
    cfg = invoice.region_config
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws.sheet_view.showGridLines = False
    fmt = f'"{cfg.currency_symbol}"#,##0.00'

    # Title
    ws["A1"] = "COMMERCIAL INVOICE"
    ws["A1"].font = Font(size=18, bold=True, color=ACCENT)
    ws.merge_cells("A1:C1")

    meta = [
        ("Invoice No.", invoice.number),
        ("Date", invoice.issue_date.strftime("%d %b %Y")),
        ("Incoterm", invoice.incoterm),
        ("Currency", cfg.currency),
    ]
    for i, (label, value) in enumerate(meta):
        ws.cell(row=1 + i, column=4, value=label).font = Font(bold=True, color=ACCENT, size=9)
        ws.cell(row=1 + i, column=5, value=value).font = Font(size=9)

    # Seller / Buyer
    ws["A3"] = "SELLER / EXPORTER"
    ws["C3"] = "BUYER / CONSIGNEE"
    for cell in ("A3", "C3"):
        ws[cell].font = Font(bold=True, color=ACCENT, size=9)

    def party_text(p):
        line = f"{p.name}\n{p.address}\n{p.city}\n{p.country}"
        if p.tax_id:
            line += f"\nTax ID: {p.tax_id}"
        return line

    ws["A4"] = party_text(invoice.seller)
    ws["C4"] = party_text(invoice.buyer)
    for cell in ("A4", "C4"):
        ws[cell].alignment = Alignment(wrap_text=True, vertical="top")
        ws[cell].font = Font(size=9)
    ws.merge_cells("A4:B4")
    ws.merge_cells("C4:E4")
    ws.row_dimensions[4].height = 70

    # Items header
    header_row = 6
    headers = ["#", "Description", "Qty", "Unit price", "Amount"]
    for col, text in enumerate(headers, 1):
        c = ws.cell(row=header_row, column=col, value=text)
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.fill = PatternFill("solid", fgColor=ACCENT)
        c.border = _border
        c.alignment = Alignment(horizontal="center")

    # Items
    r = header_row
    for i, item in enumerate(invoice.items, 1):
        r += 1
        values = [i, item.description, item.quantity, item.unit_price, item.amount]
        for col, value in enumerate(values, 1):
            c = ws.cell(row=r, column=col, value=value)
            c.border = _border
            c.font = Font(size=9)
            if i % 2 == 0:
                c.fill = PatternFill("solid", fgColor=LIGHT)
            if col in (4, 5):
                c.number_format = fmt
            if col == 1:
                c.alignment = Alignment(horizontal="center")

    # Totals
    r += 1
    totals = [("Subtotal", invoice.subtotal)]
    if cfg.show_trade_terms:
        totals += [("Freight", invoice.freight), ("Insurance", invoice.insurance)]
    totals += [("TOTAL", invoice.total)]
    for label, value in totals:
        ws.cell(row=r, column=4, value=label).alignment = Alignment(horizontal="right")
        vc = ws.cell(row=r, column=5, value=value)
        vc.number_format = fmt
        if label == "TOTAL":
            ws.cell(row=r, column=4).font = Font(bold=True, color=ACCENT)
            vc.font = Font(bold=True, color=ACCENT)
        r += 1

    # Bank details
    r += 1
    ws.cell(row=r, column=1, value="PAYMENT / BANK DETAILS").font = Font(bold=True, color=ACCENT, size=9)
    r += 1
    ws.cell(row=r, column=1, value=cfg.bank_name).font = Font(bold=True, size=9)
    for k, v in cfg.bank_details.items():
        r += 1
        ws.cell(row=r, column=1, value=f"{k}: {v}").font = Font(size=9)

    if invoice.notes:
        r += 2
        ws.cell(row=r, column=1, value="NOTES").font = Font(bold=True, color=ACCENT, size=9)
        r += 1
        nc = ws.cell(row=r, column=1, value=invoice.notes)
        nc.font = Font(size=9)
        nc.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)

    widths = {"A": 6, "B": 46, "C": 10, "D": 16, "E": 16}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    wb.save(path)
    return path
